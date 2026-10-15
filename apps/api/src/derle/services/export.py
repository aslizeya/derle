"""Export the clean table as Excel or CSV.

Excel gets real typed cells (dates as dates, money with a ₺ number format) so the file is
immediately usable; CSV is the portable, machine-friendly variant (UTF-8 with BOM so Excel
opens Turkish characters correctly, ISO dates, dot decimals).
"""

from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from derle.analytics.schema import DatasetSchema
from derle.normalize.tr.phones import mask_phone
from derle.normalize.tr.tckn import mask_tckn
from derle.storage.repo import Cursor

MONEY_FORMAT = '#,##0.00 "₺"'
DATE_FORMAT = "dd.mm.yyyy"
MONEY_ROLES = {"amount", "debit", "credit", "balance"}


def _mask(value: Any, pii: str | None) -> Any:
    if value is None or pii is None:
        return value
    text = str(value)
    if pii == "phone":
        return mask_phone(text)
    if pii == "tckn":
        return mask_tckn(text)
    if pii == "email" and "@" in text:
        user, domain = text.split("@", 1)
        return f"{user[:1]}***@{domain}"
    return value


def export_rows(
    cur: Cursor, schema: DatasetSchema, *, mask: bool
) -> tuple[list[str], list[list[Any]], list[str]]:
    """Return ``(headers, rows, roles)`` in table order, with the resolved customer next to
    the written one."""
    columns = list(schema.columns)
    select = [schema.q(c) for c in columns]
    headers = [c.label for c in columns]
    roles = [c.role for c in columns]
    pii = [c.pii for c in columns]
    if schema.customer_expr:
        position = next(i for i, c in enumerate(columns) if c.role == "party") + 1
        select.insert(position, '"party_canonical"')
        headers.insert(position, "Eşleşmiş müşteri")
        roles.insert(position, "party_canonical")
        pii.insert(position, None)
    order = schema.q(schema.date) + ", record_id" if schema.date else "record_id"
    rows = cur.execute(f"SELECT {', '.join(select)} FROM {schema.view} ORDER BY {order}").fetchall()
    out = [[_mask(v, p) if mask else v for v, p in zip(row, pii, strict=True)] for row in rows]
    return headers, out, roles


def to_xlsx(
    schema: DatasetSchema, headers: list[str], rows: list[list[Any]], roles: list[str]
) -> bytes:
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = schema.name[:31] or "Derle"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="ECE4D3")
    for row in rows:
        sheet.append([float(v) if hasattr(v, "is_finite") else v for v in row])
    for index, role in enumerate(roles, start=1):
        letter = get_column_letter(index)
        width = max(
            [len(str(headers[index - 1]))] + [len(str(r[index - 1] or "")) for r in rows[:200]]
        )
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 40)
        if role in MONEY_ROLES:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = MONEY_FORMAT
        elif role == "date":
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = DATE_FORMAT
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(
            [v.isoformat() if isinstance(v, date) else ("" if v is None else v) for v in row]
        )
    return ("﻿" + buffer.getvalue()).encode("utf-8")
