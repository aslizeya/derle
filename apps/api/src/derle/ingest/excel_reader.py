from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time

from openpyxl import load_workbook

from derle.ingest.base import NativeValue, RawGrid, RawRow, UnsupportedFileError


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    text = str(value)
    return text if text.strip() else None


def _native(value: object) -> NativeValue:
    if isinstance(value, bool | int | float | date | datetime):
        return value
    return None


class ExcelReader:
    extensions: tuple[str, ...] = (".xlsx", ".xlsm")

    def read(self, data: bytes, filename: str) -> list[RawGrid]:
        try:
            workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except (zipfile.BadZipFile, KeyError, OSError) as exc:
            raise UnsupportedFileError("corrupt_excel", str(exc)) from exc

        grids: list[RawGrid] = []
        try:
            for sheet in workbook.worksheets:
                if getattr(sheet, "sheet_state", "visible") != "visible":
                    continue
                rows: list[RawRow] = []
                for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    rows.append(
                        RawRow(
                            index=index,
                            cells=[_cell_text(v) for v in values],
                            native=[_native(v) for v in values],
                        )
                    )
                while rows and rows[-1].is_empty():
                    rows.pop()
                if sum(1 for r in rows if not r.is_empty()) >= 2:
                    grids.append(RawGrid(sheet=sheet.title, rows=rows))
        finally:
            workbook.close()

        if not grids:
            raise UnsupportedFileError("no_data_sheets")
        return grids
