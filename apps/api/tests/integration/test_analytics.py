"""Quality, dashboard, natural-language queries, history/restore and export on real data."""

import csv
import io
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from derle.analytics.dashboard import build_dashboard
from derle.analytics.quality import quality_report
from derle.analytics.schema import load_schema
from derle.config import Settings
from derle.llm.base import NullProvider
from derle.main import create_app
from derle.nlq.service import ask
from derle.services import history
from derle.services.ingest import ingest_file
from derle.storage.db import Database
from tests.conftest import make_csv

TODAY = date(2026, 10, 15)

LEDGER = make_csv(
    [
        ["Tarih", "Müşteri", "Borç", "Ödeme", "Tel"],
        ["01.05.2026", "Ali Kaya", "100", "", "0532 111 22 33"],
        ["03.05.2026", "Ali Kaya", "250", "", "0532 111 22 33"],
        ["04.05.2026", "Veli Can", "80", "", ""],
        ["10.05.2026", "Ali Kaya", "", "200", ""],
        ["02.06.2026", "Veli Can", "1.000,00", "", ""],
        ["03.06.2026", "Zeynep Ak", "40", "", ""],
        ["15/6", "zeynep ak", "60", "", ""],
        ["20.06.2026", "Veli Can", "", "500", ""],
        ["21.06.2026", "Ali Kaya", "on beş", "", ""],
    ]
)


def _setup(db: Database, workspace_id: int) -> int:
    result = ingest_file(
        db, workspace_id=workspace_id, data=LEDGER, filename="defter_2026.csv", today=TODAY
    )
    return result.tables[0].dataset_id


def _ask(db: Database, workspace_id: int, question: str) -> dict[str, Any]:
    return ask(db, NullProvider(), workspace_id=workspace_id, question=question, today=TODAY)


def test_dashboard_from_ledger(db: Database, workspace_id: int) -> None:
    dataset_id = _setup(db, workspace_id)
    with db.read() as cur:
        schema = load_schema(cur, dataset_id)
    assert schema is not None
    dash = build_dashboard(db, schema)
    kpis = {k["key"]: k["value"] for k in dash["kpis"]}
    assert kpis["records"] == 9
    assert kpis["customers"] == 3
    assert kpis["sales"] == pytest.approx(1530.0)
    assert kpis["payments"] == pytest.approx(700.0)
    assert kpis["open_balance"] == pytest.approx(830.0)
    assert kpis["period"] == ["2026-05-01", "2026-06-21"]
    charts = {c["key"]: c for c in dash["charts"]}
    assert [r["month"] for r in charts["monthly"]["rows"]] == ["2026-05", "2026-06"]
    top = charts["top_customers_open_balance"]["rows"]
    assert top[0] == {"customer": "Veli Can", "open_balance": pytest.approx(580.0)}
    assert all("SELECT" in c["sql"] for c in dash["charts"])


def test_quality_report_finds_the_unreadable_amount(db: Database, workspace_id: int) -> None:
    dataset_id = _setup(db, workspace_id)
    with db.read() as cur:
        schema = load_schema(cur, dataset_id)
        assert schema is not None
        report = quality_report(cur, schema, today=TODAY)
    codes = {i["code"]: i for i in report["issues"]}
    assert codes["unreadable_values"]["count"] == 1
    assert codes["unreadable_values"]["severity"] == "serious"
    assert codes["uncertain_values"]["params"]["year_inferred"] == 1  # "15/6"
    assert 0 < report["score"] < 100
    assert report["grade"] in {"excellent", "good", "fair", "poor"}


@pytest.mark.parametrize(
    ("question", "expected"),
    [
        ("Kaç müşterim var?", [[3]]),
        ("En çok borcu olan müşteri kim?", [["Veli Can", 580.0]]),
        ("Haziran'da ne kadar ödeme alındı?", [[500.0]]),
        ("Ali Kaya'nın borcu ne kadar?", [[150.0]]),
    ],
)
def test_questions(
    db: Database, workspace_id: int, question: str, expected: list[list[Any]]
) -> None:
    _setup(db, workspace_id)
    answer = _ask(db, workspace_id, question)
    assert answer["error"] is None, answer
    assert answer["engine"] == "rules"
    assert answer["rows"][: len(expected)] == expected
    assert answer["sql"].startswith("SELECT")


def test_relative_periods_anchor_to_stale_data(db: Database, workspace_id: int) -> None:
    _setup(db, workspace_id)
    answer = _ask(db, workspace_id, "Geçen ay en çok kim borçlandı?")
    # Data ends in June; in October "geçen ay" would be empty, so May is used and reported.
    assert answer["anchor"] == {"date": "2026-06-21", "reason": "data_end"}
    assert answer["explanation"]["period"]["start"] == "2026-05-01"
    assert answer["rows"][0] == ["Ali Kaya", 350.0]


def test_unknown_questions_and_people(db: Database, workspace_id: int) -> None:
    _setup(db, workspace_id)
    assert _ask(db, workspace_id, "Hava nasıl?")["error"] == "not_understood"
    missing = _ask(db, workspace_id, "Ahmet Demir'in borcu ne kadar?")
    assert missing["error"] == "customer_not_found"
    assert missing["rows"] == []


def test_questions_are_logged(db: Database, workspace_id: int) -> None:
    _setup(db, workspace_id)
    _ask(db, workspace_id, "Kaç müşterim var?")
    _ask(db, workspace_id, "Hava nasıl?")
    with db.read() as cur:
        rows = cur.execute("SELECT question, ok, engine FROM query_log ORDER BY id").fetchall()
    assert rows == [("Kaç müşterim var?", True, "rules"), ("Hava nasıl?", False, "none")]


def test_database_cannot_reach_the_file_system(db: Database, tmp_path: Path) -> None:
    secret = tmp_path / "secret.csv"
    secret.write_text("a,b\n1,2\n")
    with db.read() as cur, pytest.raises(Exception, match=r"(?i)external|permission|disabled"):
        cur.execute(f"SELECT * FROM read_csv('{secret.as_posix()}')").fetchall()


# --- history ----------------------------------------------------------------------------------


def _count(db: Database, dataset_id: int) -> int:
    with db.read() as cur:
        row = cur.execute(f"SELECT count(*) FROM ds_{dataset_id}").fetchone()
    assert row is not None
    return int(row[0])


def test_restore_to_before_second_upload_and_back(db: Database, workspace_id: int) -> None:
    dataset_id = _setup(db, workspace_id)
    more = make_csv(
        [["Tarih", "Müşteri", "Borç", "Ödeme", "Tel"], ["01.07.2026", "Yeni Kişi", "10", "", ""]]
    )
    ingest_file(db, workspace_id=workspace_id, data=more, filename="temmuz.csv", today=TODAY)
    assert _count(db, dataset_id) == 10
    with db.read() as cur:
        versions = history.list_versions(cur, workspace_id)
    first_upload, second_upload = versions[-1], versions[0]
    assert second_upload["added"] == 1 and second_upload["is_current"]
    assert second_upload["details"] == {"filename": "temmuz.csv"}

    with db.transaction() as cur:
        restored = history.restore(cur, workspace_id, first_upload["id"])
    assert _count(db, dataset_id) == 9

    with db.transaction() as cur:
        history.restore(cur, workspace_id, second_upload["id"])
    assert _count(db, dataset_id) == 10
    with db.read() as cur:
        latest = history.list_versions(cur, workspace_id)[0]
        statuses = dict(cur.execute("SELECT filename, status FROM uploads").fetchall())
    assert latest["kind"] == "restore"
    assert latest["details"] == {"to": second_upload["id"]}
    assert statuses == {"defter_2026.csv": "done", "temmuz.csv": "done"}
    assert restored < latest["id"]


def test_restore_brings_back_undone_values(db: Database, workspace_id: int) -> None:
    from derle.services.review import undo_transformation  # noqa: PLC0415

    dataset_id = _setup(db, workspace_id)
    with db.read() as cur:
        before = history.list_versions(cur, workspace_id)[0]["id"]
        tid, raw = cur.execute(
            "SELECT id, before_value FROM transformations WHERE kind = 'number_parsed' LIMIT 1"
        ).fetchone()  # type: ignore[misc]
    with db.transaction() as cur:
        undo_transformation(cur, tid)
    with db.read() as cur:
        undone = cur.execute(
            f"SELECT count(*) FROM ds_{dataset_id} WHERE debit IS NULL AND credit IS NULL"
        ).fetchone()
    assert undone is not None and undone[0] == 2  # the unreadable row + the undone one
    with db.transaction() as cur:
        history.restore(cur, workspace_id, before)
    with db.read() as cur:
        again = cur.execute(
            f"SELECT count(*) FROM ds_{dataset_id} WHERE debit IS NULL AND credit IS NULL"
        ).fetchone()
        flag = cur.execute(
            "SELECT undone_in_version FROM transformations WHERE id = ?", [tid]
        ).fetchone()
    assert again is not None and again[0] == 1
    assert flag == (None,)
    assert raw


def test_restore_errors(db: Database, workspace_id: int) -> None:
    _setup(db, workspace_id)
    with db.read() as cur:
        current = history.list_versions(cur, workspace_id)[0]["id"]
    with pytest.raises(history.HistoryError, match="already_current"), db.transaction() as cur:
        history.restore(cur, workspace_id, current)
    with pytest.raises(history.HistoryError, match="version_not_found"), db.transaction() as cur:
        history.restore(cur, workspace_id, 999_999)


# --- API -------------------------------------------------------------------------------------


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as test_client:
        response = test_client.post(
            "/api/v1/workspaces/1/uploads", files={"file": ("defter_2026.csv", LEDGER)}
        )
        assert response.status_code == 201
        yield test_client


def test_analytics_api(client: TestClient) -> None:
    datasets = client.get("/api/v1/workspaces/1/datasets").json()
    dataset_id = datasets[0]["id"]

    quality = client.get(f"/api/v1/datasets/{dataset_id}/quality").json()
    assert quality["score"] is not None
    assert any(i["code"] == "unreadable_values" for i in quality["issues"])

    dash = client.get(f"/api/v1/datasets/{dataset_id}/dashboard").json()
    assert {k["key"] for k in dash["kpis"]} >= {"records", "customers", "open_balance"}

    answer = client.post("/api/v1/workspaces/1/ask", json={"question": "Kaç müşterim var?"}).json()
    assert answer["rows"] == [[3]]
    assert answer["chart"] == "kpi"
    assert client.get("/api/v1/workspaces/1/questions").json()[0]["question"] == "Kaç müşterim var?"
    assert client.post("/api/v1/workspaces/1/ask", json={"question": "x"}).status_code == 422

    renamed = client.patch(f"/api/v1/datasets/{dataset_id}", json={"name": "Veresiye defteri"})
    assert renamed.json()["name"] == "Veresiye defteri"

    versions = client.get("/api/v1/workspaces/1/versions").json()
    assert versions[0]["kind"] == "upload"
    assert (
        client.post(f"/api/v1/workspaces/1/versions/{versions[0]['id']}/restore").status_code == 409
    )


def test_export_xlsx_and_csv(client: TestClient) -> None:
    dataset_id = client.get("/api/v1/workspaces/1/datasets").json()[0]["id"]

    xlsx = client.get(f"/api/v1/datasets/{dataset_id}/export", params={"format": "xlsx"})
    assert xlsx.status_code == 200
    assert "attachment" in xlsx.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    header = [c.value for c in sheet[1]]
    assert header[:3] == ["Tarih", "Müşteri", "Eşleşmiş müşteri"]
    first = [c.value for c in sheet[2]]
    assert first[0].date() == date(2026, 5, 1)
    assert sheet.cell(row=2, column=header.index("Borç") + 1).number_format.startswith("#,##0.00")

    masked = client.get(
        f"/api/v1/datasets/{dataset_id}/export", params={"format": "csv", "mask": True}
    )
    text = masked.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    tel = rows[0].index("Tel")
    assert rows[1][tel] == "0532 *** ** 33"
    assert rows[1][0] == "2026-05-01"
